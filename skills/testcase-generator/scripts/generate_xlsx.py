#!/usr/bin/env python3
"""将测试用例结构化 JSON 转换为带样式的 Excel (.xlsx)。

工作表：
  - 测试用例：标准 11 列，表头深绿底白字、冻结表头、自动筛选、优先级配色
  - 测试点清单：仅当 test_points 非空时生成

用法：
  python3 generate_xlsx.py <input.json> <output.xlsx>
"""
import sys
import os
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = ["用例编号", "所属模块/需求", "对应测试点", "测试标题", "测试类型",
           "优先级", "前置条件", "测试步骤", "测试数据", "预期结果", "备注"]
TP_HEADERS = ["测试点编号", "维度", "说明"]
COL_WIDTHS = [12, 18, 14, 26, 10, 8, 28, 40, 22, 40, 16]
TP_COL_WIDTHS = [18, 16, 70]

HEAD_FILL = PatternFill("solid", fgColor="1F4E37")  # 深绿，呼应 ProjectQuality 分类
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E37")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

PRI_FILL = {
    "P0": PatternFill("solid", fgColor="F8CBAD"),
    "P1": PatternFill("solid", fgColor="FFE699"),
    "P2": PatternFill("solid", fgColor="D9E1F2"),
    "P3": PatternFill("solid", fgColor="E2EFDA"),
}


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def build(data, out_path):
    wb = Workbook()

    # ---- 测试用例工作表 ----
    ws = wb.active
    ws.title = "测试用例"
    title = data.get("meta", {}).get("title", "测试用例集")
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(1, 1).font = TITLE_FONT
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")

    ws.append(HEADERS)
    header_row = 2
    style_header(ws, header_row, len(HEADERS))

    for tc in data.get("test_cases", []):
        row = [
            tc.get("id", ""), tc.get("module", ""), tc.get("point", ""),
            tc.get("title", ""), tc.get("type", ""), tc.get("priority", ""),
            tc.get("precondition", ""), tc.get("steps", ""), tc.get("data", ""),
            tc.get("expected", ""), tc.get("note", ""),
        ]
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
        pcell = ws.cell(r, 6)
        pkey = str(tc.get("priority", "")).strip().upper()
        if pkey in PRI_FILL:
            pcell.fill = PRI_FILL[pkey]
            pcell.alignment = Alignment(horizontal="center", vertical="top")

    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A{header_row}:{last_col}{ws.max_row}"

    # ---- 测试点清单工作表 ----
    tps = data.get("test_points", [])
    if tps:
        ws2 = wb.create_sheet("测试点清单")
        ws2.append(TP_HEADERS)
        style_header(ws2, 1, len(TP_HEADERS))
        for tp in tps:
            ws2.append([tp.get("id", ""), tp.get("dimension", ""), tp.get("desc", "")])
            r = ws2.max_row
            for c in range(1, len(TP_HEADERS) + 1):
                ws2.cell(r, c).alignment = WRAP_TOP
                ws2.cell(r, c).border = BORDER
        for i, w in enumerate(TP_COL_WIDTHS, start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return len(data.get("test_cases", [])), len(tps)


def main():
    if len(sys.argv) < 3:
        print("Usage: generate_xlsx.py <input.json> <output.xlsx>")
        sys.exit(1)
    in_path, out_path = sys.argv[1], sys.argv[2]
    with open(in_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    n_cases, n_points = build(data, out_path)
    print(f"✅ 已生成: {out_path}  (测试用例 {n_cases} 条, 测试点 {n_points} 条)")


if __name__ == "__main__":
    main()
