#!/usr/bin/env python3
"""qa-suite · 用例集导出 Excel（generate_xlsx.py）

输入 JSON schema（与 references/03-spec-model-schema.md 的用例集部分一致）：

{
  "meta": { "title": "xxx 测试用例集" },
  "test_points": [
    { "id": "TP-XX-001", "group": "分组", "desc": "描述", "category": "正向", "module": "模块名" }
  ],
  "test_cases": [
    {
      "id": "TC-XX-001", "module": "模块名", "point": "TP-XX-001,TP-XX-002",
      "title": "【分组】测试目的", "priority": "P0", "test_type": "功能",
      "precondition": "…", "steps": "1. …\\n2. …", "expected": "1. …\\n2. …", "note": "对应需求 REQ-…"
    }
  ]
}

用法：
    python3 subskills/qa-case-engine/scripts/generate_xlsx.py <input.json> <output.xlsx>

输出：含「测试用例」与「测试点清单」两个工作表，表头冻结 + 自动筛选 + 优先级配色。
"""

import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少依赖 openpyxl，请先安装：python3 -m pip install openpyxl", file=sys.stderr)
    sys.exit(1)

CASE_COLUMNS = [
    ("用例编号", "id", 16),
    ("所属模块", "module", 22),
    ("对应测试点", "point", 22),
    ("测试标题", "title", 34),
    ("优先级", "priority", 9),
    ("测试类型", "test_type", 10),
    ("前置条件", "precondition", 40),
    ("测试步骤", "steps", 52),
    ("预期结果", "expected", 52),
    ("备注", "note", 30),
]

POINT_COLUMNS = [
    ("测试点编号", "id", 16),
    ("所属模块", "module", 22),
    ("测试点类", "group", 20),
    ("类别", "category", 10),
    ("描述", "desc", 60),
]

PRIORITY_FILL = {
    "P0": PatternFill("solid", fgColor="FFC7CE"),
    "P1": PatternFill("solid", fgColor="FFEB9C"),
    "P2": PatternFill("solid", fgColor="DDEBF7"),
    "P3": PatternFill("solid", fgColor="EDEDED"),
}

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_sheet(ws, columns, rows):
    for idx, (title, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width

    for r, row in enumerate(rows, start=2):
        for c, (_, key, _) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=str(row.get(key, "") or ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if key == "priority" and cell.value in PRIORITY_FILL:
                cell.fill = PRIORITY_FILL[cell.value]
                cell.alignment = Alignment(horizontal="center", vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}"
    ws.row_dimensions[1].height = 22


def main():
    if len(sys.argv) < 3:
        print("用法：python3 generate_xlsx.py <input.json> <output.xlsx>", file=sys.stderr)
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]
    with open(in_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    cases = data.get("test_cases") or []
    points = data.get("test_points") or []
    title = (data.get("meta") or {}).get("title", "测试用例集")

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "测试用例"
    build_sheet(ws1, CASE_COLUMNS, cases)

    if points:
        ws2 = wb.create_sheet("测试点清单")
        build_sheet(ws2, POINT_COLUMNS, points)

    wb.save(out_path)
    print(f"✅ 已生成 {out_path}")
    print(f"   标题：{title}")
    print(f"   测试用例 {len(cases)} 条｜测试点 {len(points)} 个｜工作表：测试用例" + ("、测试点清单" if points else ""))


if __name__ == "__main__":
    main()